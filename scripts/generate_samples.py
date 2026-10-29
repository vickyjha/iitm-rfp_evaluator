"""Generate the synthetic supplier RFP responses used for demos and tests.

    python scripts/generate_samples.py

Writes data/sample_rfps/<category folder>/ with PDF, DOCX and DOC (Word 97-2003) proposals, the
error-case files in data/sample_rfps/error_cases/, and data/sample_rfps/sample_metadata.json.
DOC files are converted from DOCX with LibreOffice when it is installed; otherwise the committed
.doc files are kept as they are.

All companies, people, prices and references are fictional.
"""
from __future__ import annotations

import copy
import io
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sample_content_categories import (AI_GOV_RFP, AI_GOV_SUPPLIERS, DASHCAM_RFP, DASHCAM_SUPPLIERS,  # noqa: E402
                                       STATIONERY_RFP, STATIONERY_SUPPLIERS)
from sample_content_others import OTHERS_SUPPLIERS, RFP as OTHERS_RFP  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "data" / "sample_rfps"


def _mm_widths(suppliers: list[dict]) -> list[dict]:
    out = copy.deepcopy(suppliers)
    for s in out:
        for _, blocks in s["sections"]:
            for i, b in enumerate(blocks):
                if b[0] == "table" and b[2]:
                    blocks[i] = ("table", b[1], [w * mm for w in b[2]])
    return out


def _with_defaults(suppliers: list[dict], fmt: str = "pdf") -> list[dict]:
    out = copy.deepcopy(suppliers)
    for s in out:
        s["file"] = s["file"].rsplit(".", 1)[0]
        s.setdefault("format", fmt)
    return out


CATEGORY_SETS = [
    ("ai_governance", AI_GOV_RFP, _mm_widths(AI_GOV_SUPPLIERS)),
    ("stationery", STATIONERY_RFP, _mm_widths(STATIONERY_SUPPLIERS)),
    ("dashcam", DASHCAM_RFP, _mm_widths(DASHCAM_SUPPLIERS)),
    ("others", OTHERS_RFP, _with_defaults(OTHERS_SUPPLIERS)),
]
DEMO_ONLY = {"Quasar Solutions"}

styles = getSampleStyleSheet()
BODY = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.6, leading=13.2,
                      alignment=TA_LEFT, spaceAfter=5)
H1 = ParagraphStyle("h1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=19, leading=23, spaceAfter=2)
SUB = ParagraphStyle("sub", parent=BODY, fontSize=10.5, textColor=colors.HexColor("#555555"), spaceAfter=10)
H2 = ParagraphStyle("h2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12.5, leading=16,
                    spaceBefore=9, spaceAfter=4, keepWithNext=1)
CELL = ParagraphStyle("cell", parent=BODY, fontSize=8.8, leading=11.2, spaceAfter=0)
CELL_B = ParagraphStyle("cellb", parent=CELL, fontName="Helvetica-Bold")


def _tint(hex_color: str, strength: float) -> colors.Color:
    c = colors.HexColor(hex_color)
    return colors.Color(1 - (1 - c.red) * strength, 1 - (1 - c.green) * strength, 1 - (1 - c.blue) * strength)


# ============================================================================================ PDF
def build_pdf(s: dict, rfp: dict, path: Path) -> None:
    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(s["accent"]))
        canvas.rect(0, A4[1] - 9 * mm, A4[0], 9 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 8.5)
        canvas.drawString(18 * mm, A4[1] - 6 * mm, f"{s['name']}  |  Response to {rfp['reference']}")
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(18 * mm, 10 * mm, "Fictional proposal created for a classroom exercise. Confidential.")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    def table(rows, widths):
        data = [[Paragraph(str(c), CELL_B if i == 0 else CELL) for c in row] for i, row in enumerate(rows)]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _tint(s["accent"], 0.15)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=f"{s['name']} - Response to {rfp['reference']}",
                            author=s["name"])
    flow = [
        Paragraph(s["name"], ParagraphStyle("brand", parent=H1, textColor=colors.HexColor(s["accent"]))),
        Paragraph(s["tagline"], SUB),
        Paragraph(f"<b>Response to:</b> {rfp['buyer']} &mdash; {rfp['reference']}: {rfp['title']}<br/>"
                  f"<b>Submission date:</b> {s['submission_date']}", BODY),
        Spacer(1, 6),
    ]
    for heading, blocks in s["sections"]:
        flow.append(Paragraph(heading, H2))
        for block in blocks:
            if block[0] == "p":
                flow.append(Paragraph(block[1], BODY))
            elif block[0] == "bullets":
                flow.append(ListFlowable([ListItem(Paragraph(b, BODY), leftIndent=10) for b in block[1]],
                                         bulletType="bullet", start="•", leftIndent=12))
            elif block[0] == "table":
                flow += [Spacer(1, 2), table(block[1], block[2]), Spacer(1, 6)]
    doc.build(flow, onFirstPage=header_footer, onLaterPages=header_footer)


# ============================================================================================ DOCX / DOC
def build_docx(s: dict, rfp: dict, path: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    accent = RGBColor.from_string(s["accent"].lstrip("#"))
    d = docx.Document()
    d.core_properties.title = f"{s['name']} - Response to {rfp['reference']}"
    d.core_properties.author = s["name"]
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for level in (1, 2):
        d.styles[f"Heading {level}"].font.color.rgb = accent

    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"{s['name']}  |  Response to {rfp['reference']}"
    hp.runs[0].font.size = Pt(8.5)
    hp.runs[0].font.color.rgb = accent
    fp = sec.footer.paragraphs[0]
    fp.text = "Fictional proposal created for a classroom exercise. Confidential."
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.runs[0].font.size = Pt(8)

    title = d.add_heading(s["name"], level=1)
    title.runs[0].font.size = Pt(20)
    tag = d.add_paragraph(s["tagline"])
    tag.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    meta = d.add_paragraph()
    meta.add_run("Response to: ").bold = True
    meta.add_run(f"{rfp['buyer']} - {rfp['reference']}: {rfp['title']}\n")
    meta.add_run("Submission date: ").bold = True
    meta.add_run(s["submission_date"])

    for heading, blocks in s["sections"]:
        d.add_heading(heading, level=2)
        for block in blocks:
            if block[0] == "p":
                d.add_paragraph(block[1])
            elif block[0] == "bullets":
                for b in block[1]:
                    d.add_paragraph(b, style="List Bullet")
            elif block[0] == "table":
                rows = block[1]
                t = d.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, val in enumerate(row):
                        cell = t.cell(r, c)
                        cell.text = str(val)
                        run = cell.paragraphs[0].runs[0]
                        run.font.size = Pt(9.5)
                        if r == 0:
                            run.bold = True
                            shd = OxmlElement("w:shd")
                            shd.set(qn("w:val"), "clear")
                            shd.set(qn("w:fill"), s["accent"].lstrip("#") + "")
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            cell._tc.get_or_add_tcPr().append(shd)
                d.add_paragraph()
    d.save(str(path))


def docx_to_doc(docx_path: Path, doc_path: Path) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "doc", "--outdir", tmp, str(docx_path)],
                       capture_output=True, timeout=180, check=False)
        produced = Path(tmp) / (docx_path.stem + ".doc")
        if not produced.exists():
            return False
        shutil.move(str(produced), doc_path)
    return True


# ============================================================================================ error cases
def build_scanned_pdf(path: Path) -> None:
    """An image-only PDF (no text layer) that simulates a scanned proposal."""
    from PIL import Image as PILImage, ImageDraw, ImageFilter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as rl_canvas

    img = PILImage.new("L", (1240, 1754), 245)
    draw = ImageDraw.Draw(img)
    draw.text((120, 120), "ZENITH INFOTECH - PROPOSAL (SCANNED COPY)", fill=40)
    for i in range(40):
        y = 220 + i * 34
        draw.line((120, y, 120 + (900 if i % 6 else 520), y), fill=90, width=6)
    img = img.filter(ImageFilter.GaussianBlur(1.2))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    c = rl_canvas.Canvas(str(path), pagesize=A4)
    c.drawImage(ImageReader(buf), 0, 0, width=A4[0], height=A4[1])
    c.showPage()
    c.save()


# ============================================================================================ main
def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    meta: dict = {"categories": {}, "error_cases": []}
    for folder, rfp, suppliers in CATEGORY_SETS:
        target = OUT_DIR / folder
        target.mkdir(exist_ok=True)
        entries = []
        for s in suppliers:
            fmt = s.get("format", "pdf")
            if fmt == "pdf":
                path = target / f"{s['file']}.pdf"
                build_pdf(s, rfp, path)
            else:
                docx_path = target / f"{s['file']}.docx"
                build_docx(s, rfp, docx_path)
                path = docx_path
                if fmt == "doc":
                    doc_path = target / f"{s['file']}.doc"
                    if docx_to_doc(docx_path, doc_path):
                        docx_path.unlink()
                        path = doc_path
                    elif doc_path.exists():
                        docx_path.unlink()
                        path = doc_path
                        print(f"  LibreOffice not found; kept existing {doc_path.name}")
                    else:
                        print(f"  LibreOffice not found; {s['name']} kept as DOCX")
            entries.append({"file": f"{folder}/{path.name}", "supplier_name": s["name"],
                            "submission_date": s["submission_date"], "experience_rating": s["experience_rating"],
                            "format": path.suffix.lstrip("."),
                            "demo_role": "validation case: incomplete proposal" if s["name"] in DEMO_ONLY else "core"})
            print("wrote", path.relative_to(ROOT))
        meta["categories"][rfp["category_code"]] = {
            "folder": folder, "rfp_title": rfp["title"], "rfp_reference": rfp["reference"],
            "buyer": rfp["buyer"], "category_note": rfp.get("category_note", ""), "suppliers": entries,
        }

    err = OUT_DIR / "error_cases"
    err.mkdir(exist_ok=True)
    build_scanned_pdf(err / "scanned_proposal.pdf")
    # a Word document saved with a .pdf extension -> rejected by file-type validation
    shutil.copy(OUT_DIR / "stationery" / "inkwell_office_supplies.docx", err / "renamed_word_file.pdf")
    meta["error_cases"] = [
        {"file": "error_cases/scanned_proposal.pdf", "supplier_name": "Zenith Infotech",
         "submission_date": "2026-10-11", "experience_rating": 60, "format": "pdf",
         "demo_role": "error case: image-only PDF with no extractable text"},
        {"file": "error_cases/renamed_word_file.pdf", "supplier_name": "Mislabelled Supplier",
         "submission_date": "2026-10-11", "experience_rating": 50, "format": "pdf",
         "demo_role": "input validation: Word file renamed to .pdf"},
    ]
    print("wrote error_cases/")
    (OUT_DIR / "sample_metadata.json").write_text(json.dumps(meta, indent=2))
    print("wrote sample_metadata.json")


if __name__ == "__main__":
    main()
